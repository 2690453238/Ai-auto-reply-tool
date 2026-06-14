"""
导出模块
将识别和提取的结果导出为 Excel(.xlsx) 和 CSV 文件
"""

import os
import csv
from datetime import datetime
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class Exporter:
    """结果导出器"""

    @staticmethod
    def _get_default_export_path() -> str:
        """获取默认导出路径（桌面）"""
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
        if not os.path.exists(desktop):
            desktop = os.path.expanduser("~")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join(desktop, f"图片识别结果_{timestamp}")

    @staticmethod
    def export_excel(
        summaries: List[Dict[str, Any]],
        output_path: Optional[str] = None,
    ) -> str:
        """
        导出为 Excel 文件

        参数:
            summaries: 信息摘要列表（来自 InfoExtractor.batch_summarize()）
            output_path: 输出文件路径（不含扩展名），默认导出到桌面

        返回:
            实际导出的文件路径
        """
        if output_path is None:
            output_path = Exporter._get_default_export_path()
        file_path = output_path + ".xlsx"

        wb = Workbook()

        # ---- Sheet 1: 总览 ----
        ws_overview = wb.active
        ws_overview.title = "识别总览"
        Exporter._write_overview_sheet(ws_overview, summaries)

        # ---- Sheet 2: 关键字段 ----
        ws_fields = wb.create_sheet("关键字段")
        Exporter._write_fields_sheet(ws_fields, summaries)

        # ---- Sheet 3: 图片属性 ----
        ws_props = wb.create_sheet("图片属性")
        Exporter._write_properties_sheet(ws_props, summaries)

        # ---- Sheet 4: 识别全文 ----
        ws_text = wb.create_sheet("识别文字")
        Exporter._write_text_sheet(ws_text, summaries)

        # ---- Sheet 5: UI分析_候选人列表 (if available) ----
        ui_summaries = [s for s in summaries if s.get("UI分析")]
        if ui_summaries:
            ws_ui = wb.create_sheet("UI分析_候选人")
            Exporter._write_ui_candidates_sheet(ws_ui, ui_summaries)

            ws_ui_msg = wb.create_sheet("UI分析_消息")
            Exporter._write_ui_messages_sheet(ws_ui_msg, ui_summaries)

        # 删除默认创建的空白 sheet（如果存在）
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        wb.save(file_path)
        return file_path

    @staticmethod
    def export_csv(
        summaries: List[Dict[str, Any]],
        output_path: Optional[str] = None,
    ) -> str:
        """
        导出为 CSV 文件

        参数:
            summaries: 信息摘要列表
            output_path: 输出文件路径（不含扩展名），默认导出到桌面

        返回:
            实际导出的目录路径（包含多个 CSV 文件）
        """
        if output_path is None:
            output_path = Exporter._get_default_export_path()
        base_dir = output_path + "_csv"
        os.makedirs(base_dir, exist_ok=True)

        # CSV 1: 总览
        Exporter._write_csv(
            os.path.join(base_dir, "识别总览.csv"),
            summaries,
            ["文件名", "文档类型", "类型置信度", "识别状态", "文字行数", "高置信度行数", "识别文字摘要"],
            lambda s: [
                s.get("文件名", ""),
                s.get("文档类型", ""),
                s.get("类型置信度", ""),
                s.get("识别状态", ""),
                s.get("文字行数", ""),
                s.get("高置信度文字行数", ""),
                s.get("识别文字摘要", ""),
            ],
        )

        # CSV 2: 关键字段
        field_names = set()
        for s in summaries:
            for k in s.get("关键字段", {}).keys():
                field_names.add(k)

        def row_fields(s):
            row = [s.get("文件名", "")]
            fields = s.get("关键字段", {})
            for fn in sorted(field_names):
                vals = fields.get(fn, [])
                row.append(", ".join(vals) if vals else "")
            return row

        Exporter._write_csv(
            os.path.join(base_dir, "关键字段.csv"),
            summaries,
            ["文件名"] + sorted(field_names),
            row_fields,
        )

        # CSV 3: 图片属性
        prop_keys = set()
        for s in summaries:
            for k in s.get("图片属性", {}).keys():
                prop_keys.add(k)

        def row_props(s):
            row = [s.get("文件名", "")]
            props = s.get("图片属性", {})
            for pk in sorted(prop_keys):
                row.append(str(props.get(pk, "")))
            return row

        Exporter._write_csv(
            os.path.join(base_dir, "图片属性.csv"),
            summaries,
            ["文件名"] + sorted(prop_keys),
            row_props,
        )

        # CSV 4: 识别全文
        Exporter._write_csv(
            os.path.join(base_dir, "识别全文.csv"),
            summaries,
            ["文件名", "文档类型", "识别文字全文"],
            lambda s: [
                s.get("文件名", ""),
                s.get("文档类型", ""),
                s.get("识别文字全文", ""),
            ],
        )

        return base_dir

    # ========== 内部方法 ==========

    @staticmethod
    def _write_overview_sheet(ws, summaries):
        """写入总览 Sheet"""
        headers = ["序号", "文件名", "文件路径", "文档类型", "类型置信度",
                   "识别状态", "文字行数", "高置信度行数", "识别文字摘要"]
        Exporter._write_header_row(ws, headers)

        for i, s in enumerate(summaries, 1):
            row = [
                i,
                s.get("文件名", ""),
                s.get("文件路径", ""),
                s.get("文档类型", ""),
                s.get("类型置信度", ""),
                s.get("识别状态", ""),
                s.get("文字行数", ""),
                s.get("高置信度文字行数", ""),
                s.get("识别文字摘要", ""),
            ]
            ws.append(row)

        Exporter._auto_column_width(ws)

    @staticmethod
    def _write_fields_sheet(ws, summaries):
        """写入关键字段 Sheet"""
        # 收集所有字段名
        all_field_names = set()
        for s in summaries:
            for fn in s.get("关键字段", {}).keys():
                all_field_names.add(fn)
        sorted_fields = sorted(all_field_names)

        headers = ["序号", "文件名", "文档类型"] + sorted_fields
        Exporter._write_header_row(ws, headers)

        for i, s in enumerate(summaries, 1):
            row = [i, s.get("文件名", ""), s.get("文档类型", "")]
            fields = s.get("关键字段", {})
            for fn in sorted_fields:
                vals = fields.get(fn, [])
                row.append(", ".join(vals) if vals else "")
            ws.append(row)

        Exporter._auto_column_width(ws)

    @staticmethod
    def _write_properties_sheet(ws, summaries):
        """写入图片属性 Sheet"""
        # 收集所有属性键
        all_prop_keys = set()
        for s in summaries:
            for pk in s.get("图片属性", {}).keys():
                all_prop_keys.add(pk)
        sorted_props = sorted(all_prop_keys)

        headers = ["序号", "文件名"] + sorted_props
        Exporter._write_header_row(ws, headers)

        for i, s in enumerate(summaries, 1):
            row = [i, s.get("文件名", "")]
            props = s.get("图片属性", {})
            for pk in sorted_props:
                row.append(str(props.get(pk, "")))
            ws.append(row)

        Exporter._auto_column_width(ws)

    @staticmethod
    def _write_text_sheet(ws, summaries):
        """写入识别全文 Sheet"""
        headers = ["序号", "文件名", "文档类型", "识别文字全文"]
        Exporter._write_header_row(ws, headers)

        for i, s in enumerate(summaries, 1):
            row = [
                i,
                s.get("文件名", ""),
                s.get("文档类型", ""),
                s.get("识别文字全文", ""),
            ]
            ws.append(row)

        # 设置文字列宽
        ws.column_dimensions["D"].width = 80
        Exporter._auto_column_width(ws, skip_cols={"D"})

    @staticmethod
    def _write_ui_candidates_sheet(ws, ui_summaries):
        """写入 UI 分析候选人列表 Sheet"""
        headers = ["序号", "姓名", "沟通职位", "时间", "状态", "最后消息"]
        Exporter._write_header_row(ws, headers)

        row_num = 1
        for summary in ui_summaries:
            ui = summary.get("UI分析", {})
            candidates = ui.get("左侧_候选人列表", [])
            for c in candidates:
                row_num += 1
                ws.append([
                    row_num - 1,
                    c.get("name", ""),
                    c.get("job_title", ""),
                    c.get("time", ""),
                    c.get("status", ""),
                    c.get("last_message", ""),
                ])

        Exporter._auto_column_width(ws)

    @staticmethod
    def _write_ui_messages_sheet(ws, ui_summaries):
        """写入 UI 分析聊天记录 Sheet"""
        headers = ["序号", "候选人", "发送者", "消息内容"]
        Exporter._write_header_row(ws, headers)

        row_num = 1
        for summary in ui_summaries:
            ui = summary.get("UI分析", {})
            right = ui.get("右侧_详情面板", {})
            candidate_name = right.get("candidate_name", "")
            messages = right.get("messages", [])
            for m in messages:
                row_num += 1
                ws.append([
                    row_num - 1,
                    candidate_name,
                    m.get("sender", ""),
                    m.get("text", ""),
                ])

        Exporter._auto_column_width(ws)

    @staticmethod
    def _write_header_row(ws, headers: List[str]):
        """写入格式化的表头行"""
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 设置行高
        ws.row_dimensions[1].height = 25

    @staticmethod
    def _auto_column_width(ws, skip_cols: set = None):
        """自动调整列宽"""
        if skip_cols is None:
            skip_cols = set()

        for col_cells in ws.columns:
            col_letter = col_cells[0].column_letter
            if col_letter in skip_cols:
                continue

            max_length = 0
            for cell in col_cells:
                if cell.value:
                    # 中文字符算2个宽度
                    length = 0
                    for ch in str(cell.value):
                        length += 2 if ord(ch) > 127 else 1
                    max_length = max(max_length, length)

            adjusted_width = min(max_length + 4, 50)  # 最大50
            ws.column_dimensions[col_letter].width = max(adjusted_width, 8)

    @staticmethod
    def _write_csv(file_path: str, summaries: List[dict],
                   headers: List[str], row_func):
        """写入 CSV 文件"""
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for s in summaries:
                writer.writerow(row_func(s))
